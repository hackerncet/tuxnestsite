"""NCET BCA flexible enrollment and consent-based public student showcase.

Public profiles contain only explicitly published fields. Private email/mobile and
confidential HOD requests are never returned by anonymous API or HTML routes.
"""
from __future__ import annotations

import io
import json
import re
import secrets
from PIL import Image, ImageOps, UnidentifiedImageError
from datetime import datetime
from html import escape
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, DateTime, ForeignKey, String, Text, LargeBinary, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as DbSession, Mapped, mapped_column

from .database import Base, User, AuditLog, utcnow, get_db
from .student_models import BCAStudent
from .student_services import BCA_DEPARTMENT, USN_RE, can_manage_bca
from .security import current, csrf_required, hash_password
from .departments import branch_option, branch_name

router = APIRouter(tags=['NCET student enrollment and showcase'])
URL_ROOT = 'https://ncet-verifyflow.onrender.com'
MAX_BATCH = 100

class StudentShowcase(Base):
    __tablename__ = 'bca_student_showcases'
    id: Mapped[int] = mapped_column(primary_key=True)
    student_id: Mapped[int] = mapped_column(ForeignKey('users.id',ondelete='CASCADE'),unique=True,index=True)
    mobile_private: Mapped[str] = mapped_column(String(20),default='')
    email_private: Mapped[str] = mapped_column(String(180),default='')
    about: Mapped[str] = mapped_column(Text,default='')
    skills_json: Mapped[str] = mapped_column(Text,default='[]')
    certificates_json: Mapped[str] = mapped_column(Text,default='[]')
    social_links_json: Mapped[str] = mapped_column(Text,default='{}')
    published: Mapped[bool] = mapped_column(Boolean,default=False,index=True)
    photo_bytes: Mapped[bytes | None] = mapped_column(LargeBinary,nullable=True)
    photo_mime: Mapped[str | None] = mapped_column(String(32),nullable=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True),default=utcnow)

class EnrollmentItem(BaseModel):
    usn: str = Field(min_length=8,max_length=24)
    full_name: str = Field(min_length=3,max_length=130)
    branch: str = Field(min_length=2,max_length=160)
    semester: int = Field(ge=1,le=8)

class EnrollmentBody(BaseModel):
    students: list[EnrollmentItem] = Field(min_length=1,max_length=MAX_BATCH)

class FacultyUpdate(BaseModel):
    full_name: str = Field(min_length=3,max_length=130)
    branch: str = Field(min_length=2,max_length=160)
    semester: int = Field(ge=1,le=8)

class Certificate(BaseModel):
    title: str = Field(min_length=3,max_length=120)
    issuer: str = Field(min_length=2,max_length=120)
    proof_url: str = Field(min_length=12,max_length=650)

class ShowcaseEdit(BaseModel):
    full_name: str = Field(min_length=3,max_length=130)
    mobile: str = Field(default='',max_length=20)
    email: str = Field(default='',max_length=180)
    about: str = Field(default='',max_length=900)
    skills: list[str] = Field(default_factory=list,max_length=100)
    certificates: list[Certificate] = Field(default_factory=list,max_length=8)
    social_links: dict[str,str] = Field(default_factory=dict,max_length=12)
    publish: bool = False


SOCIAL_HOSTS={
 'github':'github.com','linkedin':'linkedin.com','leetcode':'leetcode.com',
 'codechef':'codechef.com','hackerrank':'hackerrank.com','kaggle':'kaggle.com',
 'codeforces':'codeforces.com','stackoverflow':'stackoverflow.com',
 'gitlab':'gitlab.com','devto':'dev.to'
}
def validate_social_links(links):
    selected={}
    for platform,raw in links.items():
        if platform not in (*SOCIAL_HOSTS,'portfolio'):
            raise HTTPException(422,'Unsupported coding profile platform.')
        link=raw.strip()
        if not link:continue
        if len(link)>400:raise HTTPException(422,'Profile link is too long.')
        url=urlsplit(link)
        host=(url.hostname or '').lower()
        if url.scheme!='https' or not host or url.username or url.password:
            raise HTTPException(422,'Profile links must use HTTPS.')
        if platform!='portfolio' and host not in (SOCIAL_HOSTS[platform],'www.'+SOCIAL_HOSTS[platform]):
            raise HTTPException(422,'Use the official '+platform+' site for that link.')
        if platform=='portfolio' and (host in ('localhost','127.0.0.1') or host.endswith(('.local','.internal'))):
            raise HTTPException(422,'Portfolio must be publicly accessible.')
        check_public_text(url.path)
        selected[platform]=link
    return selected

def valid_name(name:str)->bool:
    return 3<=len(name)<=130 and name[0] not in '=+-@' and not any(ord(c)<32 for c in name)


def validate_enrollment(db:DbSession,actor:User,items:list[EnrollmentItem])->list[tuple[str,str,str,int]]:
    result=[];seen=set()
    for i,item in enumerate(items,1):
        usn=item.usn.strip().upper()
        name=item.full_name.strip()
        if not USN_RE.fullmatch(usn):raise HTTPException(422,f'Invalid USN at entry {i}; expected 1NC followed by letters or digits.')
        if not valid_name(name):raise HTTPException(422,f'Student name is required at entry {i}.')
        if usn in seen:raise HTTPException(422,f'Duplicate USN at entry {i}.')
        branch=branch_option(db,item.branch).code
        if actor.role!='admin' and branch!=BCA_DEPARTMENT:
            raise HTTPException(403,'BCA faculty may register students only in Computer Applications.')
        seen.add(usn);result.append((usn,name,branch,item.semester))
    return result


def create_students(db, actor, items:list[EnrollmentItem]):
    can_manage_bca(actor)
    values=validate_enrollment(db,actor,items)
    if db.query(User.id).filter(User.username.in_([u.lower() for u,n,b,s in values])).first() or \
       db.query(BCAStudent.id).filter(BCAStudent.usn.in_([u for u,n,s in values])).first():
        raise HTTPException(409,'One or more USNs are already enrolled. No accounts were created.')
    creds=[]
    try:
        for usn,name,branch,semester in values:
            password=secrets.token_urlsafe(24)
            # This is deliberately not a guessable USN-derived password.
            user=User(username=usn.lower(),display_name=name,
                      role='student',department=branch,
                      password_hash=hash_password(password),active=True)
            db.add(user);db.flush()
            db.add(BCAStudent(user_id=user.id,usn=usn,semester=semester,must_change_password=True))
            creds.append({'usn':usn,'name':name,'branch':branch,'semester':semester,'temporary_password':password})
        db.add(AuditLog(actor_id=actor.id,event='bca.flexible_enrollment',detail=f'{len(values)} new USNs'))
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409,'Duplicate student account: the batch was not imported.')
    return {'created':len(creds),'credentials':creds}


@router.post('/api/bca/students/enroll')
def enroll(body:EnrollmentBody,actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    return create_students(db,actor,body.students)


@router.post('/api/bca/students/import-excel')
async def import_excel(file:UploadFile=File(...),actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    can_manage_bca(actor)
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(422,'Upload an .xlsx spreadsheet.')
    content=await file.read(350_001)
    if len(content)>350_000:raise HTTPException(413,'Spreadsheet must be smaller than 350 KB.')
    from openpyxl import load_workbook
    try:
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        sheet=book.active
        iterator=sheet.iter_rows(min_row=1,max_row=102,values_only=True)
        header=tuple(str(c or '').strip().lower() for c in next(iterator))
        aliases={'name':'full_name','department':'branch','sem':'semester'}
        normalized=tuple(aliases.get(h,h) for h in header)
        if normalized!=('usn','full_name','branch','semester'):
            raise ValueError('Columns must be USN, Name, Branch, Sem in that order.')
        parsed=[]
        for row in iterator:
            if not row or not any(c is not None and str(c).strip() for c in row):continue
            if len(parsed)>=100:raise ValueError('Maximum 100 student rows per spreadsheet.')
            if len(row)!=4 or any(c is None or not str(c).strip() for c in row):
                raise ValueError('All four fields are required for each student.')
            try:semester=int(row[3])
            except (TypeError,ValueError):raise ValueError('Semester must be a whole number between 1 and 8.')
            parsed.append(EnrollmentItem(usn=str(row[0]).strip(),
                                         full_name=str(row[1]).strip(),
                                         branch=str(row[2]).strip(),semester=semester))
    except Exception as exc:
        raise HTTPException(422,'Invalid Excel roster. Use required USN, Name, Branch, Sem (1–8).') from exc
    finally:
        if 'book' in locals():book.close()
    if not parsed:raise HTTPException(422,'The spreadsheet is empty.')
    return create_students(db,actor,parsed)


@router.patch('/api/bca/students/{usn}')
def faculty_edit(usn:str,body:FacultyUpdate,actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    can_manage_bca(actor)
    found=(db.query(User,BCAStudent).join(BCAStudent,BCAStudent.user_id==User.id)
           .filter(BCAStudent.usn==usn.upper()).first())
    if not found:raise HTTPException(404,'Student USN not found.')
    name=body.full_name.strip()
    if not valid_name(name):raise HTTPException(422,'Enter a valid student name.')
    user,student=found
    target=branch_option(db,body.branch)
    if actor.role!='admin' and (user.department!=actor.department or target.code!=user.department):
        raise HTTPException(403,'Faculty cannot transfer students between departments.')
    user.display_name=name;student.semester=body.semester;user.department=target.code
    # Changing official enrollment data requires the student to explicitly republish.
    showcase=db.query(StudentShowcase).filter(StudentShowcase.student_id==user.id).first()
    if showcase:showcase.published=False
    db.add(AuditLog(actor_id=actor.id,event='bca.student_details_updated',detail=student.usn))
    db.commit()
    return {'ok':True,'branch':user.department,'publication_review_required':bool(showcase)}


def _my_student(db,user):
    if user.role!='student':raise HTTPException(403,'Student account required.')
    found=db.query(BCAStudent).filter(BCAStudent.user_id==user.id).first()
    if not found:raise HTTPException(403,'Enrollment record missing.')
    return found


PUBLIC_EMAIL_PATTERN=re.compile(r'[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}')
PUBLIC_PHONE_PATTERN=re.compile(r'(?<!\w)\+?\d[\d() .\-]{7,}\d(?!\w)')


def check_public_text(text:str)->None:
    if PUBLIC_EMAIL_PATTERN.search(text) or PUBLIC_PHONE_PATTERN.search(text):
        raise HTTPException(422,'Remove email addresses and phone numbers from fields published publicly.')


def validate_proof(url:str)->str:
    url=url.strip()
    parts=urlsplit(url)
    if parts.scheme!='https' or not parts.hostname or parts.username or parts.password:
        raise HTTPException(422,'Certificate proof must be a public HTTPS URL.')
    if parts.hostname.lower() in ('localhost','127.0.0.1','0.0.0.0') or parts.hostname.endswith('.local'):
        raise HTTPException(422,'Invalid certificate host.')
    return url


def _showcase(db,user):
    return db.query(StudentShowcase).filter(StudentShowcase.student_id==user.id).first()


@router.get('/api/bca/showcase/me')
def my_showcase(user:User=Depends(current),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    s=_showcase(db,user)
    return {'usn':student.usn,'full_name':user.display_name if user.display_name!='Registration pending' else '',
            'department':branch_name(db,user.department),'semester':student.semester,
            'must_change_password':student.must_change_password,
            'mobile':s.mobile_private if s else '', 'email':s.email_private if s else '',
            'about':s.about if s else '', 'skills':json.loads(s.skills_json) if s else [],
            'certificates':json.loads(s.certificates_json) if s else [],
            'social_links':json.loads(s.social_links_json or '{}') if s else {},
            'published':s.published if s else False,
            'photo_available':bool(s.photo_bytes) if s else False}


@router.put('/api/bca/showcase/me')
def edit_showcase(body:ShowcaseEdit,user:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    if student.must_change_password:raise HTTPException(403,'Change your temporary password before editing your profile.')
    name=body.full_name.strip()
    if not valid_name(name):raise HTTPException(422,'Enter your complete name.')
    phone=body.mobile.strip()
    if phone and not re.fullmatch(r'\+?[0-9 ]{8,20}',phone):
        raise HTTPException(422,'Enter a valid private mobile number.')
    email=body.email.strip().lower()
    if email and (not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',email) or len(email)>180):
        raise HTTPException(422,'Enter a valid private email address.')
    skills=[]
    for item in body.skills:
        value=item.strip()
        if not 1<=len(value)<=80 or any(ord(c)<32 for c in value):
            raise HTTPException(422,'Each skill must contain 1–80 printable characters.')
        if value.casefold() not in [x.casefold() for x in skills]:skills.append(value)
    if sum(len(v) for v in skills)>5000:raise HTTPException(422,'All skills combined must be under 5,000 characters.')
    social_links=validate_social_links(body.social_links)
    check_public_text(name)
    check_public_text(body.about)
    for value in skills:check_public_text(value)
    certificates=[]
    for certificate in body.certificates:
        check_public_text(certificate.title)
        check_public_text(certificate.issuer)
        check_public_text(certificate.proof_url.split('?',1)[-1] if '?' in certificate.proof_url else '')
        certificates.append({'title':certificate.title.strip(),'issuer':certificate.issuer.strip(),
                            'proof_url':validate_proof(certificate.proof_url)})
    profile=_showcase(db,user)
    if not profile:
        profile=StudentShowcase(student_id=user.id)
        db.add(profile)
    user.display_name=name
    profile.mobile_private=phone;profile.email_private=email
    profile.about=body.about.strip()
    profile.skills_json=json.dumps(skills)
    profile.certificates_json=json.dumps(certificates)
    profile.social_links_json=json.dumps(social_links)
    profile.published=bool(body.publish)
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_updated',
                    detail=f'{student.usn}: {"public" if profile.published else "private"}'))
    db.commit()
    return {'ok':True,'published':profile.published,'public_url':f'/students/profile/{student.usn}' if profile.published else None}


PHOTO_LIMIT=150_000  # uploaded file must be strictly smaller than 150 KB (decimal)
ALLOWED_PHOTO_FORMATS={'PNG','JPEG','WEBP'}


@router.put('/api/bca/showcase/me/photo')
async def upload_profile_photo(
    file:UploadFile=File(...),
    user:User=Depends(csrf_required),
    db:DbSession=Depends(get_db),
):
    student=_my_student(db,user)
    if student.must_change_password:
        raise HTTPException(403,'Change your temporary password before uploading a photo.')
    if file.content_type not in ('image/png','image/jpeg','image/webp'):
        raise HTTPException(422,'Choose a PNG, JPEG or WebP photo under 150 KB.')
    raw=await file.read(PHOTO_LIMIT+1)
    await file.close()
    if not raw or len(raw)>=PHOTO_LIMIT:
        raise HTTPException(413,'Profile photo must be smaller than 150 KB.')
    try:
        with Image.open(io.BytesIO(raw)) as incoming:
            if incoming.format not in ALLOWED_PHOTO_FORMATS:
                raise HTTPException(422,'Unsupported photo format.')
            if incoming.width>2500 or incoming.height>2500 or incoming.width<32 or incoming.height<32:
                raise HTTPException(422,'Photo dimensions must be between 32 and 2500 pixels.')
            incoming.load()
            image=ImageOps.exif_transpose(incoming).convert('RGB')
            image.thumbnail((640,640),Image.Resampling.LANCZOS)
            photo=None
            for quality in (83,72,60,48):
                out=io.BytesIO()
                image.save(out,format='WEBP',quality=quality,method=4)
                if len(out.getvalue())<PHOTO_LIMIT:
                    photo=out.getvalue()
                    break
            if photo is None:
                raise HTTPException(422,'Photo could not fit under 150 KB. Choose a smaller image.')
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        raise HTTPException(422,'Invalid or damaged image.')
    profile=_showcase(db,user)
    if not profile:
        profile=StudentShowcase(student_id=user.id)
        db.add(profile)
    profile.photo_bytes=photo
    profile.photo_mime='image/webp'
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_photo_updated',detail=student.usn))
    db.commit()
    return {'ok':True,'photo_available':True,'size_bytes':len(photo),
            'public_photo_url':f'/students/profile/{student.usn}/photo' if profile.published else None}


@router.get('/api/bca/showcase/me/photo')
def private_profile_photo(user:User=Depends(current),db:DbSession=Depends(get_db)):
    _my_student(db,user)
    profile=_showcase(db,user)
    if not profile or not profile.photo_bytes:
        raise HTTPException(404,'Profile photo not uploaded.')
    return Response(profile.photo_bytes,media_type='image/webp',
                    headers={'Cache-Control':'private,no-store','X-Robots-Tag':'noindex,nofollow'})


@router.delete('/api/bca/showcase/me/photo')
def delete_profile_photo(user:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    profile=_showcase(db,user)
    if not profile or not profile.photo_bytes:
        raise HTTPException(404,'Profile photo not uploaded.')
    profile.photo_bytes=None
    profile.photo_mime=None
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_photo_removed',detail=student.usn))
    db.commit()
    return {'ok':True,'photo_available':False}


@router.get('/students/profile/{usn}/photo',include_in_schema=False)
def published_profile_photo(usn:str,db:DbSession=Depends(get_db)):
    if not USN_RE.fullmatch(usn.strip().upper()):
        raise HTTPException(404,'Profile unavailable.')
    row=public_query(db).filter(BCAStudent.usn==usn.strip().upper()).first()
    if not row or not row[2].photo_bytes:
        raise HTTPException(404,'Profile photo unavailable.')
    return Response(row[2].photo_bytes,media_type='image/webp',
                    headers={'Cache-Control':'no-store',
                             'X-Content-Type-Options':'nosniff'})


def public_query(db):
    return (db.query(User,BCAStudent,StudentShowcase)
            .join(BCAStudent,BCAStudent.user_id==User.id)
            .join(StudentShowcase,StudentShowcase.student_id==User.id)
            .filter(User.active.is_(True),BCAStudent.must_change_password.is_(False),
                    StudentShowcase.published.is_(True)))


def public_record(user,student,showcase):
    """Deliberate allowlist; private mobile and email never leave this function."""
    return {'name':user.display_name,'usn':student.usn,
            'department':user.department,'semester':student.semester,
            'about':showcase.about,'skills':json.loads(showcase.skills_json),
            'certificates':json.loads(showcase.certificates_json),
            'social_links':json.loads(showcase.social_links_json or '{}'),
            'photo_url':f'/students/profile/{student.usn}/photo' if showcase.photo_bytes else None,
            'updated_at':showcase.updated_at.isoformat() if showcase.updated_at else None}


@router.get('/api/bca/directory')
def public_directory(q:str='',limit:int=30,offset:int=0,db:DbSession=Depends(get_db)):
    if not 1<=limit<=50 or not 0<=offset<=1_000_000 or len(q)>60:raise HTTPException(422,'Invalid search.')
    query=public_query(db)
    if q.strip():
        term='%'+q.strip().replace('%','\\%').replace('_','\\_')+'%'
        query=query.filter(or_(BCAStudent.usn.ilike(term,escape='\\'),
                                User.display_name.ilike(term,escape='\\')))
    total=query.count()
    rows=query.order_by(BCAStudent.usn).offset(offset).limit(limit).all()
    return {'total':total,'offset':offset,'limit':limit,'items':[{'name':u.display_name,'usn':p.usn,'department':branch_name(db,u.department),
                       'skills':json.loads(s.skills_json),
                       'photo_url':f'/students/profile/{p.usn}/photo' if s.photo_bytes else None} for u,p,s in rows]}


@router.get('/api/bca/directory/{usn}')
def public_one(usn:str,db:DbSession=Depends(get_db)):
    if not USN_RE.fullmatch(usn.strip().upper()):raise HTTPException(404,'Profile unavailable.')
    record=public_query(db).filter(BCAStudent.usn==usn.strip().upper()).first()
    if not record:raise HTTPException(404,'Profile unavailable.')
    result=public_record(*record)
    result['department']=branch_name(db,record[0].department)
    return result


@router.get('/students/profile/{usn}',response_class=HTMLResponse,include_in_schema=False)
def public_profile_html(usn:str,db:DbSession=Depends(get_db)):
    info=public_one(usn,db)
    name=escape(info['name'])
    sid=escape(info['usn'])
    description=escape((info['about'] or f"{info['name']} studies at NCET.")[:155],quote=True)
    canonical=f"{URL_ROOT}/students/profile/{sid}"
    photo_url=info.get('photo_url')
    image_full=URL_ROOT+photo_url if photo_url else ''
    initials=escape(''.join(part[0] for part in info['name'].split()[:2]).upper() or 'BCA')
    photo_html=(f'<img class="profile-card-avatar" src="{escape(photo_url,quote=True)}" width="115" height="115" alt="{name} profile photo">'
                if photo_url else f'<div class="profile-card-initial" aria-label="Profile photo not added">{initials}</div>')
    skills=''.join(f'<li>{escape(x)}</li>' for x in info['skills']) or '<li>Interests not added yet</li>'
    certs=''.join('<li><strong>'+escape(x['title'])+'</strong> · '+escape(x['issuer'])+
        ' — <a rel="noopener noreferrer nofollow" target="_blank" href="'+escape(x['proof_url'],quote=True)+'">View proof ↗</a></li>'
        for x in info['certificates']) or '<li>No certificates shared yet</li>'
    socials=''.join('<li><a href="'+escape(url,quote=True)+'" rel="noopener noreferrer nofollow" target="_blank">'+escape(site.title())+' ↗</a></li>'
                    for site,url in info['social_links'].items()) or '<li>No links shared yet</li>'
    profile_schema={
        '@context':'https://schema.org','@type':'ProfilePage',
        'url':URL_ROOT+f"/students/profile/{info['usn']}",
        'dateModified':info.get('updated_at'),
        'mainEntity':{
            '@type':'Person','name':info['name'],'identifier':info['usn'],
            'description':info['about'] or f"{info['name']} is an NCET BCA student.",
            'affiliation':{'@type':'CollegeOrUniversity','name':'Nagarjuna College of Engineering and Technology'},
            'knowsAbout':info['skills'][:40],'sameAs':list(info['social_links'].values())[:10],
        },
    }
    if image_full:profile_schema['mainEntity']['image']=image_full
    schema_script=json.dumps(profile_schema,ensure_ascii=False).replace('<','\\u003c')
    og_image=(f'<meta property="og:image" content="{escape(image_full,quote=True)}">') if image_full else ''
    return HTMLResponse(f'''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#0E2455">
<title>{name} ({sid}) | NCET BCA Student Portfolio</title>
<meta name="description" content="{description}">
<meta name="robots" content="index,follow,max-image-preview:large">
<meta property="og:type" content="profile"><meta property="og:title" content="{name} | NCET BCA">
<meta property="og:description" content="{description}"><meta property="og:url" content="{canonical}">{og_image}
<script type="application/ld+json">{schema_script}</script>
<link rel="canonical" href="{canonical}"><link rel="icon" href="/static/ncet-official.webp">
<link rel="stylesheet" href="/static/bca.css"><link rel="stylesheet" href="/static/experience.css">
<link rel="stylesheet" href="/static/responsive-upgrade.css?v=20261017">
<link rel="stylesheet" href="/static/student-redesign.css?v=20261017b"></head>
<body class="student-public">
<header class="topnav"><div class="container nav-inner"><a class="brand" href="/students">
<img src="/static/ncet-official.webp" alt="NCET emblem">NCET BCA</a>
<nav class="navlinks"><a href="/students/directory">Student directory</a><a href="/students">Students</a></nav></div></header>
<main class="shell">
<div class="student-portfolio-hero">{photo_html}<div>
<span class="eyebrow">NCET · Student published portfolio</span>
<h1>{name}</h1><p>{sid} · {escape(info['department'])} · Semester {info['semester']}</p>
</div></div>
<div class="student-portfolio-columns"><div>
<section class="panel"><span class="eyebrow">ABOUT</span><h2>About me</h2>
<p>{escape(info['about']) or 'This student has not added an introduction yet.'}</p></section>
<section class="panel" style="margin-top:16px"><span class="eyebrow">SKILLS</span><h2>Skills & interests</h2>
<ul class="student-skill-chips">{skills}</ul></section></div>
<div><section class="panel"><span class="eyebrow">CONNECT</span><h2>Portfolio & coding profiles</h2>
<ul class="student-portfolio-list">{socials}</ul></section>
<section class="panel" style="margin-top:16px"><span class="eyebrow">LEARNING</span><h2>Certificates & proof</h2>
<ul class="student-portfolio-list">{certs}</ul></section></div></div>
<p class="student-portfolio-note">Shared voluntarily by the student. Linked achievements are not independently endorsed by NCET.</p>
<a class="btn" href="/students/directory">← Back to student directory</a></main>
<footer class="student-footer"><div class="container"><span>© 2026 NCET VerifyFlow · BCA</span><span><a href="/students">Student services</a> · <a href="/about">About</a></span></div></footer>
</body></html>''',headers={'Cache-Control':'no-store'})


@router.get('/students/directory',include_in_schema=False)
def directory_page(page:int=1,db:DbSession=Depends(get_db)):
    """Small paginated server-rendered lists, with every published profile crawlable."""
    from pathlib import Path
    if page<1 or page>100000:raise HTTPException(404,'Directory page not found.')
    static=Path(__file__).resolve().parents[1]/'static'
    html=(static/'student-directory.html').read_text(encoding='utf-8')
    per_page=18
    total=public_query(db).count()
    pages=max(1,(total+per_page-1)//per_page)
    if page>pages:raise HTTPException(404,'Directory page not found.')
    rows=public_query(db).order_by(BCAStudent.usn).offset((page-1)*per_page).limit(per_page).all()
    items=''.join(
        '<li><a href="/students/profile/'+escape(student.usn,quote=True)+'">'+
        escape(user.display_name)+'</a> <span class="muted">'+escape(student.usn)+'</span></li>'
        for user,student,showcase in rows
    )
    controls=(
        ('<a class="btn" href="/students/directory?page='+str(page-1)+'">← Previous</a>' if page>1 else '')+
        '<span>Page '+str(page)+' of '+str(pages)+'</span>'+
        ('<a class="btn" href="/students/directory?page='+str(page+1)+'">Next →</a>' if page<pages else '')
    )
    section=('<section class="panel server-profile-links" aria-label="Published student profiles">'+
             '<h2>Browse all published portfolios</h2>'+
             '<p class="small muted">'+str(total)+' student-published profiles. Showing '+str(len(rows))+' on this page.</p>'+
             '<ul>'+items+'</ul><nav class="profile-pagination" aria-label="Directory pages">'+controls+'</nav></section>')
    html=html.replace('</main>',section+'</main>')
    if page>1:
        html=html.replace('href="https://ncet-verifyflow.onrender.com/students/directory"',
                          'href="https://ncet-verifyflow.onrender.com/students/directory?page='+str(page)+'"')
    return HTMLResponse(html,headers={'Cache-Control':'no-store'})



def published_sitemap_paths(db:DbSession)->list[str]:
    return ['/students/directory'] + [f'/students/profile/{p.usn}'
      for p in (db.query(BCAStudent).join(User,User.id==BCAStudent.user_id)
          .join(StudentShowcase,StudentShowcase.student_id==User.id)
          .filter(User.active.is_(True),BCAStudent.must_change_password.is_(False),
                  StudentShowcase.published.is_(True)).order_by(BCAStudent.usn).limit(10000).all())]
